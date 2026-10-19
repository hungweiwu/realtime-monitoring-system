import asyncio
import io
from collections.abc import Sequence
from datetime import datetime
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.worksheet._write_only import WriteOnlyWorksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.schemas.analytics import AnalyticsParams, CategoryStat, Summary
from app.services import analytics_service
from app.services.analytics_service import MODELS, analytics_timezone, conditions

COLUMNS = {
    "records": ["id", "title", "value", "category", "timestamp", "created_by"],
    "metrics": ["value", "category", "timestamp", "is_anomaly"],
}


def _cell(ws: WriteOnlyWorksheet, value: Any, tz: ZoneInfo) -> Any:
    if isinstance(value, datetime):
        # Excel 不支援帶時區的時間：轉成分析時區後去掉 tzinfo，欄位標題註明時區
        return value.astimezone(tz).replace(tzinfo=None)
    if isinstance(value, str):
        # 強制存成文字：openpyxl 會把 "=" 開頭的字串寫成公式，
        # 使用者匯入的 title 若是 =HYPERLINK(...) 等，打開檔案時就會被執行（公式注入）
        cell = WriteOnlyCell(ws, value=value)
        cell.data_type = "s"
        return cell
    return value


async def build_workbook(session: AsyncSession, params: AnalyticsParams) -> bytes:
    settings = get_settings()
    model = MODELS[params.source]
    columns = COLUMNS[params.source]
    tz = analytics_timezone()

    # 超過上限時保留最新的資料：先依時間倒序取，再反轉成由舊到新
    stmt = (
        select(*(getattr(model, c) for c in columns))
        .where(*conditions(model, params))
        .order_by(model.timestamp.desc(), model.id.desc())
        .limit(settings.export_max_rows)
    )
    rows = list(reversed((await session.execute(stmt)).all()))
    summary = await analytics_service.summary(session, params)
    categories = await analytics_service.by_category(session, params)

    # 產生 xlsx 是純 CPU 工作（5 萬列需數秒）：放到執行緒，避免卡住 event loop 上的產生器與推送
    return await asyncio.to_thread(_render_workbook, params, columns, rows, summary, categories, tz)


def _render_workbook(
    params: AnalyticsParams,
    columns: list[str],
    rows: Sequence[Sequence[Any]],
    summary: Summary,
    categories: list[CategoryStat],
    tz: ZoneInfo,
) -> bytes:
    # write_only：逐列寫入不在記憶體保留整份工作表物件，大量資料時記憶體用量較低
    wb = Workbook(write_only=True)
    data = wb.create_sheet("data")
    data.append([f"timestamp ({tz.key})" if c == "timestamp" else c for c in columns])
    for row in rows:
        data.append([_cell(data, v, tz) for v in row])

    sheet = wb.create_sheet("summary")
    sheet.append(["source", params.source])
    sheet.append(["category", _cell(sheet, params.category or "(all)", tz)])
    sheet.append(["start", _cell(sheet, params.start, tz) if params.start else "(none)"])
    sheet.append(["end", _cell(sheet, params.end, tz) if params.end else "(none)"])
    sheet.append(["exported_rows", len(rows)])
    for key in ("count", "sum", "avg", "max", "min"):
        sheet.append([key, getattr(summary, key)])
    sheet.append([])
    sheet.append(["by category", "count", "avg", "max", "min"])
    for stat in categories:
        sheet.append([_cell(sheet, stat.category, tz), stat.count, stat.avg, stat.max, stat.min])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
