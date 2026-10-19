import io
from datetime import UTC, datetime, timedelta

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.models import Metric, Record, Role, User
from tests.conftest import auth_headers, create_user

T0 = datetime(2026, 10, 20, 0, 0, tzinfo=UTC)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def export(client: AsyncClient, user: User, **params):
    return await client.get("/api/v1/analytics/export", params=params, headers=auth_headers(user))


def workbook(resp):
    return load_workbook(io.BytesIO(resp.content))


def sheet_rows(ws) -> list[tuple]:
    return list(ws.iter_rows(values_only=True))


async def test_export_metrics_returns_xlsx_with_data_and_summary_sheets(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    user = await create_user(db_session, "alice")
    db_session.add_all(
        [
            Metric(value=1.5, category="cpu", timestamp=T0, is_anomaly=False),
            Metric(
                value=95.0, category="cpu", timestamp=T0 + timedelta(seconds=1), is_anomaly=True
            ),
        ]
    )
    await db_session.commit()

    resp = await export(client, user)

    assert resp.status_code == 200
    assert resp.headers["content-type"] == XLSX
    assert 'filename="metrics_' in resp.headers["content-disposition"]
    assert resp.headers["content-disposition"].endswith('.xlsx"')
    wb = workbook(resp)
    assert wb.sheetnames == ["data", "summary"]
    rows = sheet_rows(wb["data"])
    assert rows[0] == ("value", "category", "timestamp (Asia/Taipei)", "is_anomaly")
    assert rows[1:] == [
        (1.5, "cpu", datetime(2026, 10, 20, 8, 0), False),
        (95.0, "cpu", datetime(2026, 10, 20, 8, 0, 1), True),
    ]


async def test_export_records_applies_filters(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    user = await create_user(db_session, "alice")
    for category in ("pressure", "humidity"):
        db_session.add(
            Record(title=category, value=1, category=category, timestamp=T0, created_by=user.id)
        )
    await db_session.commit()

    resp = await export(client, user, source="records", category="pressure")

    rows = sheet_rows(workbook(resp)["data"])
    assert rows[0][:4] == ("id", "title", "value", "category")
    assert [r[1] for r in rows[1:]] == ["pressure"]


async def test_summary_sheet_contains_statistics(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    user = await create_user(db_session, "alice")
    db_session.add_all(
        [Metric(value=v, category="cpu", timestamp=T0, is_anomaly=False) for v in (2.0, 4.0)]
    )
    await db_session.commit()

    summary = dict(
        (r[0], r[1]) for r in sheet_rows(workbook(await export(client, user))["summary"]) if r[0]
    )

    assert summary["count"] == 2
    assert summary["avg"] == 3.0
    assert summary["source"] == "metrics"


async def test_text_starting_with_equals_is_written_as_string_not_formula(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    user = await create_user(db_session, "alice")
    payload = '=HYPERLINK("http://evil.example","click")'
    db_session.add(Record(title=payload, value=1, category="c", timestamp=T0, created_by=user.id))
    await db_session.commit()

    ws = workbook(await export(client, user, source="records"))["data"]

    cell = ws.cell(row=2, column=2)
    assert cell.data_type == "s"
    assert cell.value == payload


async def test_export_is_capped_and_summary_notes_truncation(
    client: AsyncClient, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(get_settings(), "export_max_rows", 2)
    user = await create_user(db_session, "alice")
    db_session.add_all(
        [
            Metric(value=float(i), category="cpu", timestamp=T0 + timedelta(seconds=i))
            for i in range(5)
        ]
    )
    await db_session.commit()

    wb = workbook(await export(client, user))

    data = sheet_rows(wb["data"])
    # 超過上限時保留最新的資料，依時間由舊到新排列
    assert [r[0] for r in data[1:]] == [3.0, 4.0]
    summary = dict((r[0], r[1]) for r in sheet_rows(wb["summary"]) if r[0])
    assert summary["exported_rows"] == 2
    assert summary["count"] == 5


async def test_viewer_can_export(client: AsyncClient, db_session: AsyncSession) -> None:
    viewer = await create_user(db_session, "watcher", role=Role.VIEWER)
    assert (await export(client, viewer)).status_code == 200


async def test_export_with_naive_start_returns_422(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    user = await create_user(db_session, "alice")
    assert (await export(client, user, start="2026-10-20T00:00:00")).status_code == 422


async def test_export_without_token_returns_401(client: AsyncClient) -> None:
    assert (await client.get("/api/v1/analytics/export")).status_code == 401


async def test_column_widths_fit_content_and_are_capped(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    user = await create_user(db_session, "alice")
    db_session.add_all(
        [
            Record(title="短", value=1, category="temperature", timestamp=T0, created_by=user.id),
            Record(title="x" * 150, value=2, category="c", timestamp=T0, created_by=user.id),
        ]
    )
    await db_session.commit()

    wb = workbook(await export(client, user, source="records"))
    data, summary = wb["data"], wb["summary"]

    # 欄位依序：id, title, value, category, timestamp, created_by
    assert data.column_dimensions["E"].width >= len("timestamp (Asia/Taipei)")
    assert data.column_dimensions["D"].width >= len("temperature")
    assert data.column_dimensions["B"].width <= 60  # 超長 title 不會把欄位撐到幾百字寬
    assert summary.column_dimensions["A"].width >= len("exported_rows")
